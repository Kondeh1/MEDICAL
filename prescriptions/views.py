"""
List prescriptions (patient sees own; provider sees issued). Create (provider only).
"""
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin
from django.http import HttpResponse
from django.shortcuts import render, redirect, get_object_or_404
from django.urls import reverse_lazy
from django.views.generic import ListView, DetailView

from audit_logs.utils import log_action
from audit_logs.models import AuditLog

from .models import Prescription, PrescriptionItem
from .forms import PrescriptionForm, PrescriptionItemFormSet


class PrescriptionListView(LoginRequiredMixin, ListView):
    """List prescriptions: patient sees received; provider sees issued."""
    model = Prescription
    template_name = "prescriptions/list.html"
    context_object_name = "prescriptions"
    paginate_by = 15

    def get_queryset(self):
        qs = Prescription.objects.select_related("patient", "provider")
        if self.request.user.is_patient:
            return qs.filter(patient=self.request.user)
        if self.request.user.is_provider:
            return qs.filter(provider=self.request.user)
        if self.request.user.is_administrator:
            return qs
        return qs.none()


class PrescriptionDetailView(LoginRequiredMixin, DetailView):
    """View prescription and items (participant or admin)."""
    model = Prescription
    template_name = "prescriptions/detail.html"
    context_object_name = "prescription"

    def get_queryset(self):
        qs = super().get_queryset().select_related("patient", "provider").prefetch_related("items")
        if self.request.user.is_administrator:
            return qs
        return qs.filter(patient=self.request.user) | qs.filter(provider=self.request.user)


def create_prescription(request):
    """Provider creates a prescription (e.g. after consultation)."""
    if not request.user.is_authenticated or not request.user.is_provider:
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied
    if hasattr(request.user, "provider_profile") and not request.user.provider_profile.is_approved:
        return redirect("accounts:provider_pending_approval")

    from django.contrib.auth import get_user_model
    User = get_user_model()
    patient_queryset = User.objects.filter(role=User.Role.PATIENT).order_by("username")

    if request.method == "POST":
        form = PrescriptionForm(request.POST)
        form.fields["patient"].queryset = patient_queryset
        formset = PrescriptionItemFormSet(request.POST)
        if form.is_valid() and formset.is_valid():
            prescription = form.save(commit=False)
            prescription.provider = request.user
            prescription.save()
            formset.instance = prescription
            formset.save()
            log_action(request, AuditLog.Action.PRESCRIPTION_CREATED, message=f"Prescription #{prescription.pk} created.")
            messages.success(request, "Prescription created.")
            return redirect("prescriptions:list")
    else:
        form = PrescriptionForm()
        form.fields["patient"].queryset = patient_queryset
        formset = PrescriptionItemFormSet()
    return render(request, "prescriptions/create.html", {"form": form, "formset": formset})


@login_required
def download_prescription(request, prescription_id, format):
    """Download prescription as PDF or DOCX."""
    from .models import Prescription
    from datetime import datetime

    prescription = get_object_or_404(Prescription, pk=prescription_id)

    # Only patient or provider can download
    if request.user != prescription.patient and request.user != prescription.provider:
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied

    if format == "pdf":
        # Generate PDF response
        response = HttpResponse(content_type="application/pdf")
        response["Content-Disposition"] = f'attachment; filename="prescription_{prescription_id}.pdf"'

        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import letter
            from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.lib.units import inch

            doc = SimpleDocTemplate(response, pagesize=letter)
            elements = []
            styles = getSampleStyleSheet()

            # Title
            title_style = ParagraphStyle(
                "CustomTitle",
                parent=styles["Heading1"],
                fontSize=24,
                textColor=colors.HexColor("#0d6efd"),
                spaceAfter=30,
                alignment=1,  # Center
            )
            elements.append(Paragraph("MEDICAL PRESCRIPTION", title_style))
            elements.append(Spacer(1, 0.2 * inch))

            # Provider and Patient Info
            info_data = [
                ["Provider:", prescription.provider.get_full_name() or prescription.provider.username],
                ["Patient:", prescription.patient.get_full_name() or prescription.patient.username],
                ["Date:", prescription.created_at.strftime("%B %d, %Y")],
            ]
            if prescription.notes:
                info_data.append(["Notes:", prescription.notes])

            info_table = Table(info_data, colWidths=[1.5 * inch, 4 * inch])
            info_table.setStyle(TableStyle([
                ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
                ('FONTNAME', (1, 0), (1, -1), 'Helvetica'),
                ('FONTSIZE', (0, 0), (-1, -1), 12),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 12),
                ('TOPPADDING', (0, 0), (-1, -1), 12),
            ]))
            elements.append(info_table)
            elements.append(Spacer(1, 0.3 * inch))

            # Medications Header
            elements.append(Paragraph("Medications:", styles["Heading2"]))
            elements.append(Spacer(1, 0.2 * inch))

            # Medications Table
            if prescription.items.exists():
                med_data = [['Medication', 'Dosage', 'Frequency', 'Duration']]
                for item in prescription.items.all():
                    med_data.append([
                        item.medication_name,
                        item.dosage,
                        item.frequency or '-',
                        item.duration or '-'
                    ])
                
                med_table = Table(med_data, colWidths=[2.5 * inch, 1.5 * inch, 1.5 * inch, 1 * inch])
                med_table.setStyle(TableStyle([
                    ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#0d6efd")),
                    ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                    ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
                    ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
                    ('FONTSIZE', (0, 0), (-1, 0), 12),
                    ('BOTTOMPADDING', (0, 0), (-1, 0), 12),
                    ('BACKGROUND', (0, 1), (-1, -1), colors.beige),
                    ('GRID', (0, 0), (-1, -1), 1, colors.black),
                ]))
                elements.append(med_table)
                
                # Instructions section
                elements.append(Spacer(1, 0.3 * inch))
                elements.append(Paragraph("Instructions:", styles["Heading3"]))
                for item in prescription.items.all():
                    if item.instructions:
                        elements.append(Paragraph(f"{item.medication_name}: {item.instructions}", styles["Normal"]))
                        elements.append(Spacer(1, 0.1 * inch))
            else:
                elements.append(Paragraph("No medications prescribed.", styles["Normal"]))

            # Build PDF
            doc.build(elements)
            return response

        except ImportError:
            # Fallback if reportlab not installed
            response = HttpResponse(content_type="text/plain")
            response["Content-Disposition"] = f'attachment; filename="prescription_{prescription_id}.txt"'
            response.write(f"Prescription #{prescription_id}\n\n")
            response.write(f"Provider: {prescription.provider.get_full_name() or prescription.provider.username}\n")
            response.write(f"Patient: {prescription.patient.get_full_name() or prescription.patient.username}\n")
            response.write(f"Date: {prescription.created_at.strftime('%B %d, %Y')}\n\n")
            if prescription.notes:
                response.write(f"Notes: {prescription.notes}\n\n")
            
            response.write("Medications:\n")
            for item in prescription.items.all():
                response.write(f"- {item.medication_name} ({item.dosage})\n")
                if item.frequency:
                    response.write(f"  Frequency: {item.frequency}\n")
                if item.duration:
                    response.write(f"  Duration: {item.duration}\n")
                if item.instructions:
                    response.write(f"  Instructions: {item.instructions}\n")
                response.write("\n")
            return response

    elif format == "docx":
        # Generate DOCX response
        try:
            from docx import Document
            from docx.shared import Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH

            doc = Document()
            
            # Title
            title = doc.add_heading('MEDICAL PRESCRIPTION', 0)
            title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            # Provider and Patient Info
            doc.add_paragraph(f'Provider: {prescription.provider.get_full_name() or prescription.provider.username}')
            doc.add_paragraph(f'Patient: {prescription.patient.get_full_name() or prescription.patient.username}')
            doc.add_paragraph(f'Date: {prescription.created_at.strftime("%B %d, %Y")}')
            if prescription.notes:
                doc.add_paragraph(f'Notes: {prescription.notes}')
            
            doc.add_paragraph()
            
            # Medications
            doc.add_heading('Medications:', level=1)
            
            if prescription.items.exists():
                table = doc.add_table(rows=1, cols=4)
                table.style = 'Table Grid'
                hdr_cells = table.rows[0].cells
                hdr_cells[0].text = 'Medication'
                hdr_cells[1].text = 'Dosage'
                hdr_cells[2].text = 'Frequency'
                hdr_cells[3].text = 'Duration'
                
                for item in prescription.items.all():
                    row_cells = table.add_row().cells
                    row_cells[0].text = item.medication_name
                    row_cells[1].text = item.dosage
                    row_cells[2].text = item.frequency or '-'
                    row_cells[3].text = item.duration or '-'
                
                # Instructions
                doc.add_paragraph()
                doc.add_heading('Instructions:', level=2)
                for item in prescription.items.all():
                    if item.instructions:
                        doc.add_paragraph(f'{item.medication_name}: {item.instructions}')
            else:
                doc.add_paragraph('No medications prescribed.')

            # Save to response
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
            response['Content-Disposition'] = f'attachment; filename=prescription_{prescription_id}.docx'
            doc.save(response)
            return response

        except ImportError:
            # Fallback if python-docx not installed
            response = HttpResponse(content_type="text/plain")
            response["Content-Disposition"] = f'attachment; filename="prescription_{prescription_id}.txt"'
            response.write(f"Prescription #{prescription_id}\n\n")
            response.write(f"Provider: {prescription.provider.get_full_name() or prescription.provider.username}\n")
            response.write(f"Patient: {prescription.patient.get_full_name() or prescription.patient.username}\n")
            response.write(f"Date: {prescription.created_at.strftime('%B %d, %Y')}\n\n")
            if prescription.notes:
                response.write(f"Notes: {prescription.notes}\n\n")
            
            response.write("Medications:\n")
            for item in prescription.items.all():
                response.write(f"- {item.medication_name} ({item.dosage})\n")
                if item.frequency:
                    response.write(f"  Frequency: {item.frequency}\n")
                if item.duration:
                    response.write(f"  Duration: {item.duration}\n")
                if item.instructions:
                    response.write(f"  Instructions: {item.instructions}\n")
                response.write("\n")
            return response

    else:
        from django.http import Http404
        raise Http404("Format not supported")
